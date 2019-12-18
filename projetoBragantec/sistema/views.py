from django.shortcuts import render, redirect, get_object_or_404
from .models import Projeto, Autor, User, Orientador, Avaliador
from django.contrib.auth import authenticate, login, get_user_model
from .forms import SubmitForm, AutorForm, RegisterForm, OrientadorForm, AvaliadorForm, EventoForm
import json
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from cruds_adminlte.crud import CRUDView
import io
from django.http import FileResponse, HttpResponse
from reportlab.pdfgen import canvas
from docx import *
from openpyxl import Workbook

User = get_user_model()

def index(request):
	return render(request, 'index.html', {})

@login_required
def dashboard(request):
	avaliadores = Avaliador.objects.filter(user_id=request.user)

	avaliadores = list(avaliadores)

	existe = False
	if len(avaliadores) is not 0:
		existe = True
	
	return render(request, 'dashboard.html', {'existe': existe})

@login_required
def register_avaliador(request):
	if request.user.is_avaliador:

		if request.method == 'POST':
			form = AvaliadorForm(request.POST)
			if form.is_valid():
				avaliador = form.save(commit=False)
				avaliador.user_id = request.user.id
				avaliador.save()
				return redirect('dashboard')
		else:
			form = AvaliadorForm()

		return render(request, 'registration/register_avaliador.html', {'form': form})
	else:
		return redirect('dashboard')

def register_user(request):
	if request.method == 'POST':
		form = RegisterForm(request.POST)
		if form.is_valid():
			user = form.save()
			user = authenticate(username=user.username, 
				password=form.cleaned_data['password1']
				)
			login(request, user)
			return redirect('dashboard')
	else:
		form = RegisterForm()
	return render(request, 'register_user.html', {'form': form})

@login_required
@csrf_exempt
def register_autor(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		return redirect('dashboard')
	else: 
		projeto = Projeto.objects.filter(user_id=request.user)
		projeto = list(projeto)

		response_data = {}

		existe_projeto = False

		if len(projeto) is not 0:
			existe_projeto = True

		if request.method == 'POST':
			data = request.POST

			email = data['email']
			nome = data['nome']
			dt_nasc = data['dt_nasc']
			curso = data['curso']
			serie = data['serie']
			instituicao = data['instituicao']

			response_data['email'] = email
			response_data['nome'] = nome
			response_data['dt_nasc'] = dt_nasc
			response_data['curso'] = curso
			response_data['serie'] = serie
			response_data['instituicao'] = instituicao

			Autor.objects.create(email=email, 
				nome=nome, 
				dt_nasc=dt_nasc, curso=curso, 
				serie=serie, instituicao=instituicao, user_id=request.user.id)

			return JsonResponse(response_data)
		else:

			form = AutorForm()
		ultimo_id = Autor.objects.latest('pk').pk

		return render(request, 'registration/register.html', {'form': form, 'existe_projeto': existe_projeto, 'ultimo_id': ultimo_id})
	

@login_required
def submit_project(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		return redirect('dashboard')
	else:
		user_id = request.user.id

		projeto = Projeto.objects.filter(user_id=request.user)
		projeto = list(projeto)
		existe_projeto = False
		if len(projeto) is not 0:
			existe_projeto = True

		orientadores = Orientador.objects.filter(user_id=user_id).values_list('email')
		autores = Autor.objects.filter(user_id = user_id).values_list('email')
		orientadores = list(orientadores)
		autores = list(autores)

		if request.method == 'POST':
			form = SubmitForm(request.POST)
			if form.is_valid():
				projeto = form.save(commit=False)
				
				if len(autores) == 1:
					projeto.email_autor_1 = autores[0][0]
					projeto.email_autor_2 = ''
					projeto.email_autor_3 = ''
				elif len(autores) == 2:
					projeto.email_autor_1 = autores[0][0]
					projeto.email_autor_2 = autores[1][0]
					projeto.email_autor_3 = ''
				elif len(autores) == 3:
					projeto.email_autor_1 = autores[0][0]
					projeto.email_autor_2 = autores[1][0]
					projeto.email_autor_3 = autores[2][0]

				if len(orientadores) == 1:
					projeto.email_orientador_1 = orientadores[0][0]
					projeto.email_orientador_2 = ''
				elif len(orientadores) == 2:
					projeto.email_orientador_1 = orientadores[0][0]
					projeto.email_orientador_2 = orientadores[1][0]

				projeto.user_id = user_id

				projeto.save()
				return redirect('confirm')
		else:
			form = SubmitForm()
		return render(request, 'registration/submit.html', {'form': form, 'existe_projeto': existe_projeto})

@login_required
def confirm(request):
	user_id = request.user.id
	orientadores = Orientador.objects.filter(user_id=user_id)
	autores = Autor.objects.filter(user_id = user_id)
	projeto = Projeto.objects.filter(user_id=user_id)
	return render(request, 'registration/confirm.html', 
		{'autores': autores,
		'projeto': projeto,
		'orientadores': orientadores})

@login_required
def register_orientador(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		return redirect('dashboard')
	else:
		projeto = Projeto.objects.filter(user_id=request.user)
		projeto = list(projeto)

		response_data = {}

		existe_projeto = False

		if len(projeto) is not 0:
			existe_projeto = True
		response_data = {}

		if request.method == 'POST':
			data = request.POST

			email = data['email']
			nome = data['nome']
			instituicao = data['instituicao']

			response_data['email'] = email
			response_data['nome'] = nome
			response_data['instituicao'] = instituicao

			Orientador.objects.create(
				email=email, 
				nome=nome, 
				instituicao=instituicao, 
				user_id=request.user.id)

			return JsonResponse(response_data)
		form = OrientadorForm()
		return render(request, 'registration/register_orientador.html', {'form': form, 'existe_projeto': existe_projeto})

@login_required
def accept_submission(request, pk):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		projeto = get_object_or_404(Projeto.objects.all(), pk=pk)
		if request.method == 'POST':
			data = request.POST
			projeto.status = data['status']
			projeto.pk = pk
			projeto.save()
			return redirect('sistema_projeto_list')
		return render(request, 'admin_actions/accept_submission.html', {'projeto':projeto})
	else:
		return redirect('dashboard')

@login_required
def reject_submission(request, pk):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		projeto = get_object_or_404(Projeto.objects.all(), pk=pk)
		if request.method == 'POST':
			data = request.POST
			projeto.status = data['status']
			projeto.pk = pk
			projeto.save()
			return redirect('sistema_projeto_list')
		return render(request, 'admin_actions/reject_submission.html', {'projeto':projeto})
	else:
		return redirect('dashboard')

@login_required
def generate_pdf(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		if request.method == 'POST':
			data = request.POST
			buffer1 = io.BytesIO()
			p = canvas.Canvas(buffer1)
			x = 720
			for campo in data:
				if campo != 'csrfmiddlewaretoken':
					p.drawString(100, x, '{}: {}'.format(campo, data[campo]))
					x-=20
				else:
					pass

			p.showPage()

			p.save()

			buffer1.seek(0)

			return FileResponse(buffer1, as_attachment=True, filename='projeto.pdf')
	else:
		return redirect('dashboard')

@login_required
def generate_docx(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		if request.method == 'POST':
			data = request.POST
			document = Document()

			for campo in data:
				if campo != 'csrfmiddlewaretoken':
					document.add_paragraph('{}: {}'.format(campo, data[campo]))
				else:
					pass

			document.add_page_break()

			response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

			response['Content-Disposition'] = 'attachment; filename=projeto.docx'

			document.save(response)

			return response
	else:
		return redirect('dashboard')

@login_required
def generate_xlsx(request):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		if request.method == 'POST':
			data = request.POST

			wb = Workbook()
			ws = wb.active

			for campo in data:
				if campo != 'csrfmiddlewaretoken':
					ws.append([campo, data[campo]])

			response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
			response['Content-Disposition'] = 'attachment; filename=projeto.xlsx'
			wb.save(response)
			return response
	else:
		return redirect('dashboard')

@login_required
def download_archive(request, pk):
	if request.user.is_staff or request.user.is_avaliador or request.user.is_diretor:
		projeto = get_object_or_404(Projeto.objects.all(), pk=pk)

		return render(request, 'admin_actions/download_archive.html', {'projeto': projeto})
	else:
		return redirect('dashboard')

@login_required
def create_event(request):
	if request.user.is_staff:
		form = EventoForm()
		return render(request, 'admin_actions/create_event.html', {'form': form})
	else:
		return redirect('dashboard')




